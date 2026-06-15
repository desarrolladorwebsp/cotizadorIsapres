#!/usr/bin/env python3
"""Parsea archivos Excel de planes Consalud (Centro, Norte, etc.) a JSON."""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import openpyxl

CLINIC_NAME_TO_ID = {
    # Compartidas / Centro
    "Centros Médicos RedSalud(A.3)": "cm-redsalud",
    "Clínica Bupa Reñaca": "cl-bupa-renaca",
    "Clínica Ciudad del Mar": "cl-ciudad-del-mar",
    "Clínica Indisa Maipú": "cl-indisa-maipu",
    "Clínica Indisa Providencia": "cl-indisa-providencia-anexo",
    "Clínicam Indisa Providencia": "cl-indisa-providencia-anexo",
    "Clínica Las Condes": "cl-las-condes",
    "Clínica Las Condes(A.3)": "cl-las-condes",
    "Clínica Los Carrera": "cl-los-carrera",
    "Clínica Los Leones": "cl-los-leones",
    "Clínica Meds": "cl-meds",
    "Clínica Meds(A.3)": "cl-meds",
    "Clínica RedSalud Rancagua": "cl-redsalud-rancagua",
    "Clínica RedSalud Valparaiso": "cl-redsalud-valparaiso",
    "Clínica RedSalud Valparaíso": "cl-redsalud-valparaiso",
    "Clínica RedSalud Vitacura": "cl-redsalud-vitacura",
    "Clínica RedSalud Providencia": "cl-redsalud-providencia",
    "Clínica RedSalud Santiago": "cl-redsalud-santiago",
    "Clínica San Carlos de Apoquindo": "cl-san-carlos",
    "Clínica Santa Maria": "cl-santa-maria",
    "Clínica Santa María(A.3)": "cl-santa-maria",
    "Clínica Santa María (A.3)": "cl-santa-maria",
    "Clínica Universidad de los Andes": "cl-univ-andes",
    "Hospital Clinico Fusat": "hosp-clinico-fusat",
    "Hospital Clinico de Viña del Mar": "hosp-vina-del-mar",
    "Hospital Clínico Universidad Católica": "hosp-clinico-uc",
    "Integramédica": "integramedica",
    "Red de Salud UC Christus": "red-uc-christus",
    "Red de Salud UC Christus(A.3)": "red-uc-christus-a3",
    # Norte
    "Clínica San José de Arica": "cl-san-jose-arica",
    "Clínica Atacama": "cl-atacama",
    "Clínica RedSalud Iquique": "cl-redsalud-iquique",
    "Clínica La Portada": "cl-la-portada",
    "Clínica La Portad": "cl-la-portad",
    "Clínica Tarapacá": "cl-tarapaca",
    "Clínica Andes Salud El Loa": "cl-andes-salud-el-loa",
    "Clínica RedSalud Elqui": "cl-redsalud-elqui",
    "Clínica Bupa Antofagasta": "cl-bupa-antofagasta",
    "Clínica Alemana de Santiago(A.)": "cl-alemana-santiago-a",
    "Clínica Alemana de Santiago(A.2)": "cl-alemana-santiago-a2",
    "Clínica Alemana de Santiago(A.3)": "cl-alemana-santiago-a3",
    "Clínica Dávila": "cl-davila",
    "Clínica Dávila(A.3)": "cl-davila-a3",
    "Clínica Dávila(A.3": "cl-davila-a3",
    "Clínica Dávila Vespucio": "cl-davila-vespucio",
    "Clínica Bupa Santiago": "cl-bupa-santiago",
    "Clínica  Indisa Providencia": "cl-indisa-providencia-anexo",
    "Clínica RedSalud Providencia,": "cl-redsalud-providencia",
    "Clínica Med": "cl-meds",
    "Hospital Clínico Universidad Católica(A.2)": "hosp-clinico-uc-a2",
    "Hospital Clínico Universidad de Chile": "hosp-clinico-uch",
    "Hospital Clínico Universidad de Chile(A.2)": "hosp-clinico-uch-a2",
    # Octava
    "Clínica Los Andes Los Ángeles": "cl-los-andes-la",
    "Clínica Los Andes Los Ángele": "cl-los-andes-la",
    "Clínica Andes Salud Talca": "cl-andes-salud-talca",
    "Clínica Andes Salud Chillán": "cl-andes-salud-chillan",
    "Clínica Andes Salud Chillá": "cl-andes-salud-chillan",
    "Clínica Andes Salud Chillán,": "cl-andes-salud-chillan",
    "Clínica Andes Salud Concepción": "cl-andes-salud-concepcion",
    "Clínica Biobío": "cl-biobio",
    "Sanatorio Alemán": "sanatorio-aleman",
    "Clínica Indisa Maipú,": "cl-indisa-maipu",
    "Clínica Meds(A.3),": "cl-meds",
    "Clínica RedSaludm Santiago": "cl-redsalud-santiago",
    # Santiago
    "Centros Médicos RedSalud (A.1)": "cm-redsalud",
    "Centros Médicos RedSalud (A.3)": "cm-redsalud",
    "Centros Médicos RedSalud (A.4)": "cm-redsalud",
    "Centros Médicos RedSalud(A.3": "cm-redsalud",
    "Clínica Alemana de Santiago(A.3),": "cl-alemana-santiago-a3",
    "Clínica Bupan Santiago": "cl-bupa-santiago",
    "Clínica Cordillera": "cl-cordillera",
    "Clínica Dávila (A.3)": "cl-davila-a3",
    "Clínica Dávila Vespucio,": "cl-davila-vespucio",
    "Clínica DávilaVespucio": "cl-davila-vespucio",
    "Clínica Hospital del Profesor": "cl-hospital-del-profesor",
    "Clínica Indisa": "cl-indisa-providencia-anexo",
    "Clínica RedSalud Santiago,": "cl-redsalud-santiago",
    "Clínica San Carlos de Apoquindo,": "cl-san-carlos",
    "Clínica Santa María(A.": "cl-santa-maria",
    "Hospital Clínico Universidad de Chile (A.2)": "hosp-clinico-uch-a2",
    "Hospital Clínico Universidad de Chile (A.3)": "hosp-clinico-uch-a3",
    "Hospital Clínico Universidad de Chile (A.4)": "hosp-clinico-uch-a4",
    "Hospital Parroquial de San Bernardo": "hosp-parroquial-san-bernardo",
    "Red de Salud UC Christus (A.3)": "red-uc-christus-a3",
    "Red de Salud UC Christus (A.4)": "red-uc-christus-a4",
}

HOSPITALARIO_LABELS = {"PORCENTAJE HOSPITALARIO", "HOSPITALARIO"}
AMBULATORIO_LABELS = {"PORCENTAJE AMBULATORIO", "AMBULATORIO"}
COVERAGE_HEADER_PATTERN = re.compile(r"(\d+)%\s*Sin Tope", re.IGNORECASE)


def parse_price(value) -> float | None:
    if value is None:
        return None
    return float(str(value).strip().replace(",", "."))


def normalize_cell_text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.startswith('"') and not text.endswith('"'):
        text = text[1:]
    return text.strip('"')


def parse_coverage_text(text: str, coverage_type: str) -> list[dict]:
    if not text:
        return []

    entries: list[dict] = []
    matches = list(COVERAGE_HEADER_PATTERN.finditer(text))

    for index, match in enumerate(matches):
        percentage = int(match.group(1))
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        section = text[start:end]

        for line in section.split("\n"):
            clinic_name = line.strip()
            if not clinic_name:
                continue

            clinic_id = CLINIC_NAME_TO_ID.get(clinic_name)
            if not clinic_id:
                print(f"Clínica sin mapeo (omitida): {clinic_name!r}", file=sys.stderr)
                continue

            entries.append(
                {
                    "clinic_id": clinic_id,
                    "clinic_name": clinic_name,
                    "percentage": percentage,
                    "type": coverage_type,
                }
            )

    return entries


def parse_workbook(xlsx_path: Path) -> list[dict]:
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
    worksheet = workbook["Hoja 1"]
    plans: list[dict] = []

    for row_idx in range(1, worksheet.max_row + 1):
        if worksheet.cell(row_idx, 1).value != "CODIGO":
            continue

        codes: list[str] = []
        columns: list[int] = []
        for col in range(2, worksheet.max_column + 1):
            value = worksheet.cell(row_idx, col).value
            if value:
                codes.append(str(value).strip())
                columns.append(col)

        if not codes:
            continue

        prices: dict[str, float | None] = {}
        hospitalario: dict[str, str] = {}
        ambulatorio: dict[str, str] = {}

        for row in range(row_idx + 1, min(row_idx + 12, worksheet.max_row + 1)):
            label = worksheet.cell(row, 1).value
            if label == "CODIGO":
                break
            if not label:
                continue

            for code, col in zip(codes, columns):
                cell_value = worksheet.cell(row, col).value
                if label == "PRECIO BASE":
                    prices[code] = parse_price(cell_value)
                elif label in HOSPITALARIO_LABELS and cell_value:
                    hospitalario[code] = normalize_cell_text(cell_value)
                elif label in AMBULATORIO_LABELS and cell_value:
                    ambulatorio[code] = normalize_cell_text(cell_value)

        for code in codes:
            coverage: list[dict] = []
            hosp_text = hospitalario.get(code)
            amb_text = ambulatorio.get(code)

            if hosp_text:
                coverage.extend(parse_coverage_text(hosp_text, "hospitalaria"))
            if amb_text:
                coverage.extend(parse_coverage_text(amb_text, "ambulatoria"))

            plans.append(
                {
                    "isapre": "Consalud",
                    "plan_name": f"Consalud {code}",
                    "unique_code": code,
                    "base_price_uf": prices.get(code),
                    "has_top": False,
                    "additional_notes": None,
                    "pdf_url": None,
                    "pdf_public_id": None,
                    "coverage": coverage,
                }
            )

    return plans


def main() -> None:
    if len(sys.argv) < 3:
        print(
            "Uso: parse-consalud-excel.py <ruta-al-xlsx> <ruta-salida-json>",
            file=sys.stderr,
        )
        sys.exit(1)

    xlsx_path = Path(sys.argv[1]).expanduser().resolve()
    output_path = Path(sys.argv[2]).expanduser().resolve()
    plans = parse_workbook(xlsx_path)
    output_path.write_text(json.dumps(plans, ensure_ascii=False, indent=2), encoding="utf-8")

    with_cov = sum(1 for plan in plans if plan["coverage"])
    print(f"Planes parseados: {len(plans)} (con cobertura: {with_cov}) -> {output_path}")


if __name__ == "__main__":
    main()
